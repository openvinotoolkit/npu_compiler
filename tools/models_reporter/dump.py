#
# Copyright (C) 2023 Intel Corporation
# SPDX-License-Identifier: Apache-2.0
#

from dataclasses import dataclass
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
from openpyxl.formula.translate import Translator
from table import DataTable
import utils


@dataclass
class ExcelSheet:
    name: str
    data_frame: pd.DataFrame
    format: bool
    header: bool
    index: bool


class ExcelReport:
    tables: dict = {}

    def append_df(self, name: str, data_frame: pd.DataFrame, format=True, header=True, index=False):
        if name not in self.tables:
            excel_sheet = ExcelSheet(name, data_frame, format, header, index)
            self.tables[name] = excel_sheet
        else:
            df = self.tables[name].data_frame
            self.tables[name].data_frame = pd.concat([df, data_frame])

    def append(self, table: DataTable, format=True, header=True, index=False):
        self.append_df(table.name, table.to_data_frame(),
                       format, header, index)


def adjust_column_width(df, sheet, format=None):
    max_column_width = 50

    for column in df:
        column_width = max(df[column].astype(str)
                           .map(len).max(), len(str(column)))
        column_width = 3 + min(column_width, max_column_width)
        col_idx = df.columns.get_loc(column)
        sheet.set_column(col_idx, col_idx, column_width, format)


def report_filename(device):
    now = datetime.now()
    time_string = now.strftime("%Y_%m_%d_%Hh%Mm")
    year, month, day, time = time_string.split("_")
    work_week, work_day = utils.date_to_work_week(
        int(year), int(month), int(day))
    return f"{device}_ww{work_week:02d}.{work_day}_{time}.xlsx"


def read_excel_template():
    workbook = load_workbook(filename='models_report_template.xlsx')
    statistics = workbook["Statistics"]
    compilation_problems = workbook["CompilationProblems"]
    inference_problems = workbook["InferenceProblems"]
    error_tracking = workbook["ErrorTracking"]

    statistics_df = pd.DataFrame(statistics.values)
    compilation_problems_df = pd.DataFrame(compilation_problems.values)
    inference_problems_df = pd.DataFrame(inference_problems.values)
    tracking_df = pd.DataFrame(error_tracking.values)
    return statistics_df, compilation_problems_df, inference_problems_df, tracking_df


def color_red_failed_models(df, sheet, red_bkg):
    for i, row in enumerate(df.itertuples(index=False)):
        if getattr(row, "status") == "FAILED":
            sheet.set_row(i + 1, None, red_bkg)


def add_autofilter(df, sheet):
    sheet.autofilter(0, 0, df.shape[0], df.shape[1] - 1)


def fill_problems_table(writer, df, sheet, problem_frequency):
    sorted_problems = sorted(problem_frequency.items(),
                             key=lambda kv: (kv[1], kv[0]), reverse=True)

    # Translate all formulas
    formulas = df.loc[2]
    for row in range(len(sorted_problems) - 1):
        for col in range(len(formulas)):
            formula = formulas[col]
            translated = Translator(
                formula, "A1").translate_formula("A" + str(2 + row))
            sheet.write(row + 3, col, translated)

    regexp_row_height = 45
    sheet.set_row(1, regexp_row_height)

    error_format = writer.book.add_format()
    error_format.set_text_wrap()

    # Populate unique error messages and its count
    for row, (err, count) in enumerate(sorted_problems):
        sheet.write(row + 2, 2, count)
        sheet.write(row + 2, 3, err, error_format)

    # Wrap regexp cells
    regexp_format = writer.book.add_format()
    regexp_format.set_text_wrap()
    for col_num in range(4, len(df.columns.values) - 4):
        sheet.write(1, col_num, None, regexp_format)

    # Adjust width of columns
    sheet.set_column(1, 1, 13)
    sheet.set_column(3, 3, 70)

    # Filter
    sheet.autofilter(1, 2, df.shape[0] - 1, df.shape[1] - 2)


def adjust_tracking_table(worksheet):
    worksheet.set_column(0, 0, 11)
    regexp_row_height = 90
    worksheet.set_row(0, regexp_row_height)


def count_unique_problems(compilation_df):
    problems = {}
    for row in compilation_df.itertuples(index=False):
        err = getattr(row, "errorId")
        if not err:
            continue
        problems[err] = problems.get(err, 0) + 1

    return problems


def export_table(writer, name, table):
    table.data_frame.to_excel(writer, sheet_name=name,
                              index=table.index, header=table.header)
    if table.format:
        add_autofilter(table.data_frame, writer.sheets[name])
        adjust_column_width(
            table.data_frame, writer.sheets[name])


def to_excel(report: ExcelReport, report_path):
    tab_order = ["Report", "Package", "Statistics", "Compilation", "CompilationProblems",
                 "Inference", "InferenceProblems", "UnsupportedLayer", "ErrorTracking"]

    print(f"Saving to {report_path}...")
    with pd.ExcelWriter(report_path) as writer:
        for name in tab_order:
            if name in report.tables:
                export_table(writer, name, report.tables[name])

        for name in report.tables:
            if name not in tab_order:
                export_table(writer, name, report.tables[name])

        # Fill CompilationProblems table
        compilation_df = report.tables["Compilation"].data_frame
        compilation_problems_df = report.tables["CompilationProblems"].data_frame
        compilation_problems_sheet = writer.sheets["CompilationProblems"]
        problem_frequency = count_unique_problems(compilation_df)
        fill_problems_table(writer, compilation_problems_df,
                            compilation_problems_sheet, problem_frequency)

        red_style = writer.book.add_format()
        red_style.set_bg_color('#FFC7CE')
        red_style.set_font_color('#9C0006')
        color_red_failed_models(
            compilation_df, writer.sheets["Compilation"], red_style)

        # Fill InferenceProblems table
        if "Inference" in report.tables:
            inference_df = report.tables["Inference"].data_frame
            inference_problems_df = report.tables["InferenceProblems"].data_frame
            inference_problems_sheet = writer.sheets["InferenceProblems"]
            problem_frequency = count_unique_problems(inference_df)
            fill_problems_table(writer, inference_problems_df,
                                inference_problems_sheet, problem_frequency)

            color_red_failed_models(
                inference_df, writer.sheets["Inference"], red_style)

        adjust_tracking_table(writer.sheets["ErrorTracking"])

    print("Done, heh")
